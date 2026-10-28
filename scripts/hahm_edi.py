#!/usr/bin/env python
"""HAHM — ตรวจ EDI: MANIFEST × ENTER (ไฟล์ ENTER ชุดรวมหลาย forwarder)

ไฟล์ ENTER ของ HAHM เป็น PDF ชุดรวม: ARRIVAL NOTICE/สำเนา B/L ของ Heung-A (+ ATTACHED RIDER),
ใบ MARK&NO./CNTR/B/L NO., ตาราง AMENDMENT (Cargoport / ASIA ATLANTIC / INITIAL / Handle ฯลฯ),
จดหมายแก้ไขแบบข้อความอิสระ, หน้าภาพสแกน และฟอร์มของ forwarder ที่ skill มีตัวอ่านอยู่แล้ว
(MAC-NELS, THE M CORP/CMS, DSV)

สคริปต์นี้แยกหน้า ENTER ตามชนิดฟอร์ม แล้วอ่านแต่ละชนิดด้วยตัวอ่านของตัวเอง — ฟอร์มที่ skill
manifest-enter-edi-check อ่านได้อยู่แล้วจะส่งเฉพาะหน้านั้นให้ build_edi.py ของ skill อ่าน (เรียกใช้
อย่างเดียว ไม่แก้ไขไฟล์ของ skill) จากนั้นใช้ parse_manifest()/build_report() ของ skill สร้าง EDI.xlsx
รูปแบบเดิม แล้วเติมเลขหน้า ENTER ของแต่ละ B/L และหมายเหตุหน้าภาพสแกนลงคอลัมน์หมายเหตุ

ใช้:  python hahm_edi.py                      (หาไฟล์ใน input/ ให้อัตโนมัติ, ผลอยู่ใน output/EDI.xlsx)
      python hahm_edi.py --manifest input/MANIFEST-HAHM.xls --enter input/ENTER-HAHM.pdf --shed 0141
"""
import argparse
import glob
import importlib.util
import os
import re
import sys
import tempfile

import openpyxl
import pymupdf
from openpyxl.styles import PatternFill

# build_edi.py: ใช้สำเนาที่วางข้างไฟล์นี้ก่อน (ชุดที่แจกเพื่อน) ไม่มีจึงใช้ของ skill manifest-enter-edi-check
_HERE = os.path.dirname(os.path.abspath(__file__))
SKILL_SCRIPT = next((p for p in (os.path.join(_HERE, 'build_edi.py'),
                                 os.path.join(os.path.expanduser('~'), '.claude', 'skills', 'manifest-enter-edi-check',
                                              'scripts', 'build_edi.py')) if os.path.exists(p)), '')


def _load_shared():
    if not SKILL_SCRIPT:
        sys.exit('!! ไม่พบ build_edi.py (วางไว้โฟลเดอร์เดียวกับ hahm_edi.py)')
    spec = importlib.util.spec_from_file_location('build_edi_shared', SKILL_SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


be = _load_shared()

# เลขตู้ ISO 6346: ตัวที่ 4 เป็น U/J/Z เสมอ — กันเลข SHIPMENT ID ('SSGN0000001') และส่วนท้ายของเลข B/L
CONT2_RE = re.compile(r'(?<![A-Z])[A-Z]{3}[UJZ]\d{7}')
BL_TXT_RE = re.compile(r'HASL[A-Z]\d{8,}[A-Z]?')
BL_IN_TXT_RE = re.compile(r'HASL[A-Z]\d{8,11}[A-Z]?(?![A-Za-z0-9])')   # ในข้อความที่มีช่องว่าง
STATUS_KW = re.compile(r'STATUS|ลากตู้|ลากตต|เปิดตู้|ขนส่ง|\bCFS\b', re.I)
PKG_WORD_RE = re.compile(r'^(?:WOODEN\s+)?(PALLETS?|PLTS?|CARTONS?|CTNS?|CASES?|BOX(?:ES)?|BOX\(ES\)|PACKAGES?|PKGS?|'
                         r'SETS?|ROLLS?|REELS?|DRUMS?|BAGS?|BALES?|CRATES?|UNITS?|PIECES?|BUNDLES?|SKIDS?)\b', re.I)
REEFER_RE = re.compile(r'(?:MINUS|-)\s*(\d{1,2}(?:\.\d)?)\s*(?:DEG|DEGREES|°|\'\s*C)', re.I)


# ═══════════════ ตัวช่วยอ่านหน้า PDF ═══════════════
def page_lines(pg):
    out = []
    for b in pg.get_text('dict')['blocks']:
        for l in b.get('lines', []):
            t = ''.join(s['text'] for s in l['spans']).strip()
            if t:
                out.append({'x': l['bbox'][0], 'y': l['bbox'][1], 'x1': l['bbox'][2], 't': t})
    return out


def page_annots(pg):
    out = []
    for a in pg.annots() or []:
        if a.type[1] != 'FreeText':
            continue
        txt = a.info.get('content', '') or ''
        parts = [re.sub(r'\s+', ' ', p).strip() for p in re.split(r'[\r\n]+', txt)]
        parts = [p for p in parts if p]
        if parts:
            out.append({'rect': a.rect, 'lines': parts, 't': ' '.join(parts)})
    return out


def nospace(s):
    return re.sub(r'\s+', '', s).upper()


def in_annot(l, annots):
    """บรรทัดนี้เป็นข้อความที่ FreeText annotation วาดไว้บนหน้า (อยู่ในกรอบของมัน)"""
    return any(a['rect'].x0 - 3 <= l['x'] <= a['rect'].x1 and a['rect'].y0 - 3 <= l['y'] <= a['rect'].y1 + 1
               for a in annots)


def conts_in(*txts):
    seen = []
    for t in txts:
        for c in CONT2_RE.findall(str(t or '').upper()):
            if c not in seen:
                seen.append(c)
    return seen


def split_pkg(s):
    """'38 BUNDLES' / '600BAGS' / '1 PALLET(S)' / '222 PLTS' -> (จำนวน, ชนิด)"""
    m = re.match(r'^\s*(\d[\d,]*)\s*([A-Za-z].*)?$', s or '')
    if not m:
        return None, ''
    return int(m.group(1).replace(',', '')), (m.group(2) or '').strip()


def clean_status(parts):
    out = []
    for p in parts:
        p = p.replace('ลากตต', 'ลากตู้')
        p = re.sub(r'^\s*(?:STATUS|REMARKS?)\s*[:=]?\s*', '', p, flags=re.I).strip(' *"“”:')
        if p and p not in out:
            out.append(p)
    return ' '.join(out)


def is_status_line(t):
    return bool(STATUS_KW.search(t))


def find_reefer(*txts):
    for t in txts:                      # 'MINUS 18 DEGREES' / 'TEMP : - 18'C' ก่อน — find_temp อ่านแบบมีช่องว่างไม่ออก
        m = REEFER_RE.search(t or '')
        if m:
            return f'-{m.group(1)}°C'
    return be.find_temp(*txts)


def company_name(lines):
    """ชื่อบริษัทจากบล็อกชื่อ+ที่อยู่: เอาบรรทัดจนเจอคำลงท้ายบริษัท (ไม่เกิน 3 บรรทัด)"""
    lines = [l for l in lines if l.strip()]
    if not lines:
        return ''
    out = []
    for l in lines[:3]:
        out.append(l.strip())
        if re.search(r'\b(LTD|LIMITED|CORP(?:ORATION)?|INC|PLC|LLC|CO\.?,?\s*LTD)\b\.?\s*\*?$', l.strip(), re.I) \
                or re.search(r'\b(LTD|LIMITED)\b', l, re.I):
            break
    else:
        out = [lines[0].strip()]
    s = ' '.join(out)
    s = re.sub(r'\s+', ' ', s.replace(' D.', 'D.')).strip(' *')
    return re.sub(r'LT D\.', 'LTD.', s)


def blank(bl, page):
    return {'bl': bl, 'cons': '', 'cont': '', 'pkgs': None, 'pkgtype': '', 'gw': None, 'meas': None,
            'marks': '', 'desc': '', 'status_raw': '', 'transit_raw': '', 'dg': '', 'reefer': '',
            'low_confidence': False, 'pages': [page]}


def finish(rec, *extra_txt):
    blob = ' '.join([rec['marks'], rec['desc']] + [t for t in extra_txt if t])
    rec['dg'] = rec['dg'] or be.find_dg(blob)
    rec['reefer'] = rec['reefer'] or find_reefer(blob)
    if not rec['transit_raw'] and be.TRANSIT_RE.search(blob):
        rec['transit_raw'] = blob
    return rec


# ═══════════════ จำแนกชนิดหน้า ═══════════════
def page_kind(pg, lines):
    ns = nospace(' '.join(l['t'] for l in lines))
    if not ns:
        return 'SCAN'
    if 'MAC-NELS' in ns and 'ATTACHEDSHEETS' in ns:
        return 'SHARED_MACNELS'
    if 'INBOUNDCARGOMANIFEST' in ns:
        return 'SHARED_CMS'
    if 'AMENDMENTOFMANIFEST' in ns and 'B/LCHANGED' in ns:
        return 'SHARED_DSV'
    if 'ATTACHEDRIDER' in ns and "SHIPPER'SLOAD" not in ns:
        return 'RIDER'
    if 'MARK&NO.' in ns and 'CNTR:' in ns and 'B/LNO.:' in ns:
        return 'MARKNO'
    if "SHIPPER'SLOAD" in ns:
        return 'BLFORM'
    up = [l['t'].upper() for l in lines]
    if any(u in ('CONSIGNEE',) for u in up) and any(u.startswith('QUANTITY') for u in up) \
            and any(u.startswith('DESCRIPTION') for u in up) and any(u.startswith('GROSS WEIGHT') for u in up):
        return 'TABLE'
    if len(lines) <= 3 and not BL_TXT_RE.search(ns):
        return 'SCAN'       # ภาพสแกนที่มีแค่ FreeText แปะทับ (เช่นเลขตู้)
    return 'LETTER'


# ═══════════════ ARRIVAL NOTICE / สำเนา B/L ของ Heung-A ═══════════════
def _lab(lines, key):
    for l in lines:
        if nospace(l['t']).startswith(key):
            return l
    return None


def parse_blform(lines, annots, page):
    bl_line = next((l for l in lines if BL_TXT_RE.fullmatch(nospace(l['t'])) and l['x'] > 330 and l['y'] < 130), None)
    if not bl_line:
        return {}
    rec = blank(nospace(bl_line['t']), page)
    rec['src'] = 'BLFORM'
    body = [l for l in lines if not in_annot(l, annots)]

    sl = _lab(body, "SHIPPER'SLOAD")
    y_sl, x_sl = sl['y'], sl['x']
    tots = [l['y'] for l in (_lab(body, '(*)TOTALNO'), _lab(body, 'SAY:')) if l]
    y_t = min(tots) if tots else y_sl + 230
    kgs = next((l for l in body if abs(l['y'] - y_sl) < 4 and re.search(r'KGS', l['t'])), None)
    cbm = next((l for l in body if abs(l['y'] - y_sl) < 4 and re.search(r'CBM', l['t'])), None)
    pkl = next((l for l in body if abs(l['y'] - y_sl) < 4 and l['x'] < x_sl - 2 and re.match(r'^\d', l['t'])), None)
    x_gw = (kgs['x'] if kgs else 420) - 6
    rec['gw'] = be.num(kgs['t']) if kgs else None
    rec['meas'] = be.num(cbm['t']) if cbm else None
    if pkl:
        rec['pkgs'], rec['pkgtype'] = split_pkg(pkl['t'])

    # ผู้รับ / ผู้รับแจ้ง: ใช้ป้าย Consignee / Notify Party / Pre Carriage ถ้ามี ไม่งั้นแบ่งบล็อกจากช่องว่างแนวตั้ง
    lc, ln, lp = _lab(body, 'CONSIGNEE'), _lab(body, 'NOTIFYPARTY'), _lab(body, 'PRECARRIAGE')
    left = sorted((l for l in body if l['x'] < 200 and l['y'] < y_sl - 30 and not BL_TXT_RE.search(nospace(l['t']))),
                  key=lambda l: l['y'])
    if lc and ln:
        y_end = lp['y'] if lp else ln['y'] + 100
        cons_lines = [l['t'] for l in left if lc['y'] + 3 < l['y'] < ln['y'] - 1]
        noti_lines = [l['t'] for l in left if ln['y'] + 3 < l['y'] < y_end - 1]
    else:
        blocks, cur = [], []
        for l in left:
            if cur and l['y'] - cur[-1]['y'] > 16:
                blocks.append(cur)
                cur = []
            cur.append(l)
        if cur:
            blocks.append(cur)
        cons_lines = [l['t'] for l in blocks[1]] if len(blocks) > 1 else []
        noti_lines = [l['t'] for l in blocks[2]] if len(blocks) > 2 else []
    cons = company_name(cons_lines)
    notify = company_name(noti_lines)
    if re.match(r'^TO\s+ORDER', cons, re.I) and notify and not re.match(r'^SAME\s+AS', notify, re.I):
        rec['cons_note'] = f'ENTER Consignee = {cons} → ใช้ชื่อ Notify'
        cons = notify
    rec['cons'] = cons
    if notify and not re.match(r'^SAME\s+AS', notify, re.I):
        rec['notify'] = notify

    # พื้นที่สินค้า: ใต้แถว SHIPPER'S LOAD ถึงบรรทัด (*) Total No.
    cargo = [l for l in body if y_sl + 3 < l['y'] < y_t - 1]
    cont_line = lambda t: bool(CONT2_RE.search(t.upper())) and bool(
        re.match(r"^(?:R?\d{2}[A-Z]{0,3}'?)?\s*[A-Z]{3}[UJZ]\d{7}", nospace(t)))
    noise = lambda t: bool(re.match(r"^\d{2}'?[A-Z]*X\d+$", nospace(t))) or nospace(t).startswith('PARTICULARSFURNISHED') \
        or nospace(t) in ("SAIDTOCONTAIN", "SHIPPER'SLOAD,COUNT&WEIGHT")
    marks, desc, typed = [], [], []
    for l in sorted(cargo, key=lambda l: (l['y'], l['x'])):
        t = l['t']
        if noise(t):
            continue
        if is_status_line(t) and t != 'CY/CY':
            typed.append(t)
            continue
        if l['x'] < x_sl - 10:
            if not cont_line(t):
                marks.append(t)
        elif l['x'] < x_gw:
            desc.append(t)
    rec['marks'] = ' '.join(marks)
    rec['desc'] = ' '.join(desc)
    rec['cont'] = ';'.join(conts_in(*[l['t'] for l in cargo if l['x'] < x_sl - 10 and cont_line(l['t'])],
                                    *[a['t'] for a in annots]))

    # STATUS ที่ลูกค้าพิมพ์เพิ่ม (ทั้งข้อความแบนราบบนหน้า และ FreeText annotation)
    for l in body:
        if l['y'] >= y_t - 1 or l['y'] <= y_sl + 3:
            if is_status_line(l['t']) and l['t'] != 'CY/CY':
                typed.append(l['t'])
    labelled = _apply_labelled_annots(rec, annots, lc, ln)
    for a in annots:
        if a['t'] in labelled:
            continue
        typed += [p for p in a['lines'] if is_status_line(p)]
        rec['dg'] = rec['dg'] or be.find_dg(a['t'])
        rec['reefer'] = rec['reefer'] or find_reefer(a['t'])
        if be.TRANSIT_RE.search(a['t']):
            rec['transit_raw'] = a['t']
    rec['status_raw'] = rec['status_raw'] or clean_status(typed)
    return {rec['bl']: finish(rec)}


def _apply_labelled_annots(rec, annots, lc, ln):
    """FreeText ที่ลูกค้าพิมพ์ทับ: ช่อง Consignee (แก้ชื่อผู้รับ) หรือข้อความมีป้าย 'QUANTITY PACKAGE:' ฯลฯ"""
    used = set()
    LBL = [('pkgs', r'(?:QUANTITY\s*PACKAGE|NO\.?\s*OF\s*PKGS?|PACKAGES?)'), ('desc', r'(?:CARGO\s*DESCRIPTION|DESCRIPTION)'),
           ('marks', r'MARKS?\s*&\s*NOS?\.?'), ('gw', r'GROSS\s*WEIGHT'), ('meas', r'(?:MEAS(?:UREMENT)?|CBM)'),
           ('status', r'STATUS')]
    for a in annots:
        hits = {}
        free = []
        for p in a['lines']:
            for k, pat in LBL:
                m = re.match(rf'^{pat}\s*[:=]\s*(.*)$', p, re.I)
                if m:
                    hits[k] = m.group(1).strip()
                    break
            else:
                free.append(p)
        in_cons_box = lc is not None and ln is not None and lc['y'] - 10 <= a['rect'].y0 < ln['y'] - 2
        if hits or in_cons_box:
            used.add(a['t'])
            # ชื่อผู้รับ = บรรทัดแรกที่ไม่มีป้าย เฉพาะเมื่อพิมพ์ทับช่อง Consignee หรือเป็นใบแก้ไขหลายช่อง (≥2 ป้าย)
            if free and (in_cons_box or len(hits) >= 2) and not CONT2_RE.search(free[0].upper()):
                rec['cons'] = free[0]
            if 'pkgs' in hits:
                rec['pkgs'], rec['pkgtype'] = split_pkg(hits['pkgs'])
            for k in ('desc', 'marks'):
                if k in hits:
                    rec[k] = hits[k]
            if 'gw' in hits:
                rec['gw'] = be.num(hits['gw'])
            if 'meas' in hits:
                rec['meas'] = be.num(hits['meas'])
            if 'status' in hits:
                rec['status_raw'] = clean_status([hits['status']])
    return used


def parse_rider(lines, page):
    """ATTACHED RIDER ของ B/L: คืน (bl, marks, typed status lines)"""
    bl = next((nospace(l['t']) for l in lines if BL_TXT_RE.fullmatch(nospace(l['t']))), '')
    mi = next((l for l in lines if 'MARKINFORMATION' in nospace(l['t'])), None)
    marks, typed = [], []
    for l in sorted(lines, key=lambda l: (l['y'], l['x'])):
        if is_status_line(l['t']):
            typed.append(l['t'])
            continue
        if mi and l['y'] > mi['y'] + 3 and l['x'] < 250:
            marks.append(l['t'])
    return bl, ' '.join(marks), typed


# ═══════════════ ใบ MARK&NO. / NO. OF PKGS / DESCRIPTION OF GOODS (2 B/L ต่อหน้า) ═══════════════
def parse_markno(lines, page):
    out = {}
    lbls = sorted((l for l in lines if nospace(l['t']) == 'B/LNO.:'), key=lambda l: l['y'])
    hdr = next((l for l in lines if nospace(l['t']) == 'MARK&NO.'), None)
    prev_y = (hdr['y'] + 5) if hdr else 0
    for lb in lbls:
        val = next((l for l in lines if abs(l['y'] - lb['y']) < 3 and l['x'] > lb['x'] + 20
                    and BL_TXT_RE.fullmatch(nospace(l['t']))), None)
        if not val:
            continue
        rec = blank(nospace(val['t']), page)
        blk = [l for l in lines if prev_y < l['y'] <= lb['y'] + 3]
        prev_y = lb['y'] + 3
        cl = next((l for l in blk if nospace(l['t']) == 'CNTR:'), None)
        y_cntr = cl['y'] if cl else lb['y'] - 20
        skip_y = [y_cntr, lb['y']]
        marks, desc, mid = [], [], []
        for l in sorted(blk, key=lambda l: (l['y'], l['x'])):
            t = l['t']
            if t.startswith('G.W.'):
                rec['gw'] = be.num(t)
            elif t.startswith('MEAS.'):
                rec['meas'] = be.num(t)
            elif any(abs(l['y'] - y) < 3 for y in skip_y) and l['x'] < 200:
                if abs(l['y'] - y_cntr) < 3:
                    rec['cont'] = ';'.join(conts_in(t)) or rec['cont']
            elif l['x'] < 200:
                marks.append(t)
            elif l['x'] < 340:
                mid.append(l)
            elif l['y'] < y_cntr:
                desc.append(t)
        cnt_l = next((l for l in mid if re.fullmatch(r'[\d,]+', l['t'])), None)
        if cnt_l:
            rec['pkgs'] = int(cnt_l['t'].replace(',', ''))
            pt = next((l for l in mid if l is not cnt_l and l['y'] < cnt_l['y'] + 30 and not re.fullmatch(r'[\d,]+', l['t'])), None)
            rec['pkgtype'] = pt['t'] if pt else ''
            rec['cons'] = ' '.join(l['t'] for l in mid if l is not cnt_l and l is not pt and l['y'] < y_cntr)
        rec['marks'] = ' '.join(marks)
        rec['desc'] = ' '.join(desc)
        out[rec['bl']] = finish(rec)
    return out


# ═══════════════ ตาราง AMENDMENT (Cargoport / ASIA ATLANTIC / INITIAL / Handle ฯลฯ) ═══════════════
def _status_value(lines, l):
    m = re.match(r'^\s*STATUS\s*[:=]?\s*(.*)$', l['t'], re.I)
    v = (m.group(1) if m else l['t']).strip(' :')
    if not v:
        v = ' '.join(o['t'] for o in sorted(lines, key=lambda o: o['x'])
                     if o is not l and abs(o['y'] - l['y']) < 3 and o['x'] > l['x'] and o['t'].strip(' :'))
    return clean_status([v])


def parse_table_page(lines, annots, page):
    """คืน (rows:dict[bl->rec], page_status, page_conts)"""
    hdr = next(l for l in lines if l['t'].upper() == 'CONSIGNEE')
    hy = hdr['y']
    H = [l for l in lines if abs(l['y'] - hy) < 6]
    col = {}
    for l in H:
        u = l['t'].upper()
        if u.startswith('MARKS'):
            col['MARKS'] = l['x']
        elif u.startswith('QUANTITY'):
            col['QTY'] = l['x']
        elif u.startswith('DESCRIPTION'):
            col['DESC'] = l['x']
        elif u == 'CONSIGNEE':
            col['CONS'] = l['x']
        elif u.startswith('GROSS WEIGHT'):
            col['GW'] = l['x']
        elif u.startswith('HB/L'):
            col['HBL'] = l['x']
    bounds = sorted((x - 45, k) for k, x in col.items())

    def col_of(x):
        name = 'BL'
        for left, k in bounds:
            if x >= left:
                name = k
        return name

    tot = [l['y'] for l in lines if l['y'] > hy + 5 and (l['t'].upper() == 'TOTAL' or l['t'].upper().startswith('TOTAL '))]
    pg_end = [l['y'] for l in lines if re.match(r'^Page \d+ of \d+', l['t']) or l['t'].startswith('[FM-')]
    end_y = min(tot + pg_end + [1e9])

    anchors = []
    used = set()
    for l in sorted(lines, key=lambda l: l['y']):
        if l['y'] > hy + 5 and l['y'] < end_y and col_of(l['x']) == 'BL' and re.match(r'^HASL[A-Z0-9]+$', l['t']):
            bl = l['t']
            suf = next((o for o in lines if o is not l and abs(o['x'] - l['x']) < 5 and 3 < o['y'] - l['y'] < 16
                        and re.fullmatch(r'[A-Z0-9]{1,2}', o['t'])), None)
            if suf:
                bl += suf['t']
                used.add(id(suf))
            used.add(id(l))
            anchors.append((l['y'], bl))
    rows = {}
    for i, (ay, bl) in enumerate(anchors):
        ny = anchors[i + 1][0] if i + 1 < len(anchors) else end_y
        rl = [l for l in lines if ay - 4 <= l['y'] < ny - 4 and id(l) not in used]
        rec = blank(bl, page)
        parts = {k: [] for k in ('BL', 'MARKS', 'QTY', 'DESC', 'CONS', 'GW', 'HBL')}
        st = []
        for l in sorted(rl, key=lambda l: (l['y'], l['x'])):
            c = col_of(l['x'])
            if re.match(r'^\s*STATUS\b', l['t'], re.I):
                st.append(_status_value(rl, l))
                continue
            parts[c].append(l['t'])
        rec['marks'] = ' '.join(parts['MARKS'])
        rec['desc'] = ' '.join(parts['DESC'])
        rec['cons'] = ' '.join(parts['CONS'])
        rec['pkgs'], rec['pkgtype'] = split_pkg(' '.join(parts['QTY']))
        gwt = ' '.join(parts['GW'])
        mg = re.search(r'G\.?\s*W\.?\s*([\d,]+\.?\d*)', gwt) or re.search(r'([\d,]+\.\d+)\s*KGS', gwt)
        mm = re.search(r'([\d,]+\.?\d*)\s*(?:M3|CBM)', gwt)
        mm_v = None
        if not mm:   # '25.0000' / 'M3' แยกบรรทัด
            nums = [x for x in re.findall(r'[\d,]+\.\d+', gwt)]
            mm_v = be.num(nums[1]) if len(nums) > 1 else None
        rec['gw'] = be.num(mg.group(1)) if mg else None
        rec['meas'] = be.num(mm.group(1)) if mm else mm_v
        rec['cont'] = ';'.join(conts_in(*parts['BL'], *parts['MARKS'], *parts['DESC']))
        rec['status_raw'] = ' '.join(st)
        rows[bl] = rec
    page_status = [_status_value(lines, l) for l in lines
                   if re.match(r'^\s*STATUS\b', l['t'], re.I) and not any(ay - 4 <= l['y'] < end_y for ay, _ in anchors[:1])]
    page_conts = conts_in(*[l['t'] for l in lines if l['y'] < hy], *[a['t'] for a in annots])
    for a in annots:
        if is_status_line(a['t']):
            page_status.append(clean_status([a['t']]))
    return rows, [s for s in page_status if s], page_conts


# ═══════════════ จดหมายแก้ไขแบบข้อความอิสระ ═══════════════
_LETTER_STOP = re.compile(r'^(CONSIGNEE|CNEE|CONTAINER|CONT\.?\s*NO|TOTAL|-{5,}|WE |YOUR|THANK|STATUS|REMARK|'
                          r'NOTIFY|PLACE OF|VESSEL|จ่าย)', re.I)
_L_CONS = r'(?:CONSIGNEE\s*NAME|CONSIGNEE|CNEE)'
_L_CONT = r'(?:CONTAINER\s*NO\.?|CONT(?:R)?\.?\s*(?:NO\.?)?)'
_L_PKG = r'(?:PACKAGES?|NO\.\s*OF\s*PCS|QUANTITY\s*PACKAGE)'
_L_GW = r'(?:WEIGHT\s*&\s*CBM|GROSS\s*(?:WEIGHT|WT\.?))'
_L_MEAS = r'(?:GROSS\s*CBM|CBM\.?)'
_L_MARKS = r'(?:MARKS?\s*(?:AND|&)\s*(?:NUMBERS?|NOS?\.?|NO\.?))'
_L_DESC = r'(?:CARGO\s*DESCRIPTION|DESCRIPTION)'


def _label_value(lines, pat):
    for l in sorted(lines, key=lambda l: l['y']):
        m = re.match(rf'^{pat}\s*[:=]?\s*(.*)$', l['t'], re.I)
        if not m:
            continue
        v = m.group(1).strip(' :=')
        if not v:
            right = sorted((o for o in lines if o is not l and abs(o['y'] - l['y']) < 6 and o['x'] > l['x']),
                           key=lambda o: o['x'])
            v = ' '.join(o['t'] for o in right).strip(' :=')
        if not v:
            near = sorted((o for o in lines if o['x'] > l['x'] + 40 and o['t'].startswith(':')
                           and abs(o['y'] - l['y']) < 22), key=lambda o: abs(o['y'] - l['y']))
            if near:
                v = near[0]['t'].strip(' :=')
        if v:
            return v, l
    return '', None


def parse_letter(lines, annots, page):
    body = [l for l in lines if not in_annot(l, annots)]
    txt = ' '.join(l['t'] for l in body)
    bls = BL_IN_TXT_RE.findall(txt)
    if not bls:
        return {}
    rec = blank(bls[0], page)

    cons, _ = _label_value(body, _L_CONS)
    rec['cons'] = cons
    rec['as_per_attached'] = bool(re.search(r'AS\s*PER\s*ATTACHED', cons, re.I))
    rec['cont'] = ';'.join(conts_in(txt, *[a['t'] for a in annots]))

    # ตารางที่มีหัว MARKS ... DESCRIPTION บนแถวเดียวกัน
    # หัวตารางจริงไม่มี ':' (ต่างจากป้าย 'Marks and Number : ABC' ที่มีค่าต่อท้าย)
    hm = next((l for l in body if re.match(_L_MARKS, l['t'], re.I) and ':' not in l['t'] and
               any(re.search(r'DESCRIP', o['t'], re.I) and ':' not in o['t'] and abs(o['y'] - l['y']) < 8
                   for o in body)), None)
    if hm:
        hrow = [o for o in body if abs(o['y'] - hm['y']) < 10]
        hx = {}
        for o in hrow:
            u = o['t'].upper()
            if re.match(_L_MARKS, u):
                hx.setdefault('MARKS', o['x'])
            if re.search(r'NUMBER$|NO\.?\s*OF\s*P|PKGS|PACKAGES', u) and not re.match(_L_MARKS, u):
                hx.setdefault('PKG', o['x'])
            if 'DESCRIP' in u and not re.match(_L_MARKS, u):
                hx.setdefault('DESC', o['x'])
            if 'WEIGHT' in u and not re.match(_L_MARKS, u):
                hx.setdefault('GW', o['x'])
        single = len(hrow) == 1          # หัวตารางพิมพ์เป็นข้อความบรรทัดเดียว (ไม่รู้ตำแหน่งคอลัมน์)
        rows, last_y = [], max(o['y'] for o in hrow)
        for l in sorted((o for o in body if o['y'] > last_y + 3), key=lambda o: (o['y'], o['x'])):
            if rows and l['y'] - max(r['y'] for r in rows) > 45:
                break
            if _LETTER_STOP.match(l['t']) or is_status_line(l['t']):
                if rows:
                    break
                continue
            rows.append(l)
        marks, desc, gwm, nums_only = [], [], [], []
        pkg_l = next((r for r in rows if PKG_WORD_RE.match(split_pkg(r['t'])[1] or '') or
                      (re.fullmatch(r'\d[\d,]*', r['t']) and 'PKG' in hx and abs(r['x'] - hx['PKG']) < 40)), None)
        if pkg_l:
            rec['pkgs'], rec['pkgtype'] = split_pkg(pkg_l['t'])
            if not rec['pkgtype']:
                pw = next((r for r in rows if PKG_WORD_RE.fullmatch(r['t'].strip()) and abs(r['x'] - pkg_l['x']) < 40), None)
                rec['pkgtype'] = pw['t'] if pw else ''
        pkg_x = pkg_l['x'] if pkg_l else hx.get('PKG', 200)
        desc_x = hx.get('DESC')
        gw_x = hx.get('GW')
        for r in rows:
            t = r['t']
            if r is pkg_l or (PKG_WORD_RE.fullmatch(t.strip()) and abs(r['x'] - pkg_x) < 40):
                continue
            if re.fullmatch(r'[\d,]+\.\d+', t):
                nums_only.append(r)
            elif re.fullmatch(r'(KGS?|CBM|M3)\.?', t.strip(), re.I):
                continue                              # หน่วยแยกบรรทัด — จับคู่กับตัวเลขด้านล่าง
            elif re.search(r'KGS?\b|CBM|M3\b', t, re.I) or (gw_x and r['x'] >= gw_x - 30):
                gwm.append(t)
            elif r['x'] < pkg_x - 5:
                marks.append(t)
            elif single or desc_x is None or r['x'] >= (pkg_x + desc_x) / 2:
                desc.append(t)
        rec['marks'] = ' '.join(marks)
        rec['desc'] = ' '.join(desc)
        # ตัวเลขล้วน ('10,774.40') กับหน่วยที่พิมพ์บรรทัดล่าง ('KGS') — จับคู่ตามคอลัมน์ (x ใกล้กัน)
        for r in nums_only:
            unit = next((u for u in rows if re.fullmatch(r'(KGS?|CBM|M3)\.?', u['t'].strip(), re.I)
                         and 0 < u['y'] - r['y'] < 30 and abs(u['x'] - r['x']) < 25), None)
            gwm.append(f"{r['t']} {unit['t'] if unit else ''}".strip())
        g = ' '.join(gwm)
        mg = re.search(r'([\d,]+\.?\d*)\s*KGS?', g, re.I)
        mm = re.search(r'([\d,]+\.?\d*)\s*(?:CBM|M3)', g, re.I) or re.search(r'CBM\.?\s*=?\s*([\d,]+\.?\d*)', g, re.I)
        nums = re.findall(r'\d[\d,]*\.\d+', g)
        rec['gw'] = be.num(mg.group(1)) if mg else (be.num(nums[0]) if nums else None)
        rec['meas'] = be.num(mm.group(1)) if mm else (be.num(nums[1]) if len(nums) > 1 else None)
        # น้ำหนัก/ปริมาตรที่พิมพ์แบบมีป้ายใต้ตาราง ('GROSS WT. = …' / 'CBM. = …')
        if rec['gw'] is None:
            v, _ = _label_value(body, _L_GW)
            rec['gw'] = be.num(v) if v else None
        if rec['meas'] is None:
            v, _ = _label_value(body, _L_MEAS)
            rec['meas'] = be.num(v) if v else None
    else:
        v, _ = _label_value(body, _L_PKG)
        if v:
            rec['pkgs'], rec['pkgtype'] = split_pkg(v)
        v, _ = _label_value(body, _L_GW)
        if v:
            rec['gw'] = be.num(v)
            mm = re.search(r'([\d,]+\.?\d*)\s*(?:M3|CBM)', v, re.I)
            if mm:
                rec['meas'] = be.num(mm.group(1))
        if rec['meas'] is None:
            v, _ = _label_value(body, _L_MEAS)
            rec['meas'] = be.num(v) if v else None
        for key, pat in (('marks', _L_MARKS), ('desc', _L_DESC)):
            v, l = _label_value(body, pat)
            if l is not None:
                vx = l['x'] + 60
                more = []
                for o in sorted((o for o in body if l['y'] + 5 < o['y'] < l['y'] + 60 and o['x'] >= vx),
                                key=lambda o: o['y']):
                    if _LETTER_STOP.match(o['t']) or re.match(r'^[A-Z .&\'/-]+\s*:', o['t']) or o['t'].startswith(':') \
                            or any(p['x'] < vx and abs(p['y'] - o['y']) < 4 for p in body):
                        break                 # ค่าของป้ายถัดไป (มีป้ายอยู่ทางซ้ายแถวเดียวกัน)
                    more.append(o['t'])
                rec[key] = ' '.join([v] + more).strip()
    typed = [l['t'] for l in body if is_status_line(l['t']) and not re.match(_L_DESC, l['t'], re.I)]
    rec['status_raw'] = clean_status(typed + [p for a in annots for p in a['lines'] if is_status_line(p)])
    rec['combined_text'] = True           # ข้อความอิสระ — MARKS/DESC ไม่ตรงให้เป็นสีส้ม (ให้คนยืนยัน)
    for a in annots:
        rec['dg'] = rec['dg'] or be.find_dg(a['t'])
        rec['reefer'] = rec['reefer'] or find_reefer(a['t'])
    return {rec['bl']: finish(rec, *[a['t'] for a in annots])}


# ═══════════════ ฟอร์มที่ skill อ่านได้อยู่แล้ว: ส่งเฉพาะหน้านั้นให้ build_edi.py ═══════════════
def parse_shared(doc, pages, bl_re):
    sub = pymupdf.open()
    for p in pages:
        sub.insert_pdf(doc, from_page=p - 1, to_page=p - 1, annots=True)
    fd, tmp = tempfile.mkstemp(suffix='.pdf')
    os.close(fd)
    try:
        sub.save(tmp)
        sub.close()
        ent, declared = be.parse_enter(tmp, bl_re)
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass
    page_bls = {}
    for p in pages:
        for bl in BL_IN_TXT_RE.findall(doc[p - 1].get_text()):
            page_bls.setdefault(bl, []).append(p)
    out = {}
    for bl, rec in ent.items():
        if bl in page_bls:
            rec['pages'] = sorted(set(page_bls[bl]))
            out[bl] = rec
    return out, declared


# ═══════════════ อ่าน ENTER ทั้งชุด ═══════════════
def parse_enter_hahm(path, bl_re):
    doc = pymupdf.open(path)
    kinds, L, A = [], [], []
    for i in range(doc.page_count):
        pg = doc[i]
        ls, an = page_lines(pg), page_annots(pg)
        L.append(ls)
        A.append(an)
        kinds.append(page_kind(pg, ls))

    ENT, LETTERS, scans = {}, {}, []
    info = {'kinds': kinds, 'scans': scans, 'groups': []}

    def put(recs):
        for bl, rec in recs.items():
            if bl in ENT:
                ENT[bl]['pages'] = sorted(set(ENT[bl]['pages'] + rec['pages']))
                for k, v in rec.items():
                    if k != 'pages' and v not in (None, '') and ENT[bl].get(k) in (None, ''):
                        ENT[bl][k] = v
            else:
                ENT[bl] = rec

    # กลุ่มหน้าที่ติดกันชนิดเดียวกัน
    groups, i = [], 0
    while i < len(kinds):
        j = i
        while j + 1 < len(kinds) and kinds[j + 1] == kinds[i] and kinds[i].startswith('SHARED'):
            j += 1
        groups.append((kinds[i], list(range(i + 1, j + 2))))
        i = j + 1
    info['groups'] = groups

    table_doc = None                      # ตาราง AMENDMENT หลายหน้า: STATUS/เลขตู้ระบุครั้งเดียวทั้งฉบับ
    table_docs = []
    for kind, pages in groups:
        p = pages[0]
        ls, an = L[p - 1], A[p - 1]
        if kind.startswith('SHARED'):
            recs, _ = parse_shared(doc, pages, bl_re)
            if kind == 'SHARED_DSV':
                _fix_dsv(recs, doc)
            # FreeText STATUS ที่หัวใบแรก (เหนือ B/L แรก) = คำสั่งของทั้งชุด ไม่ใช่แค่ B/L แรก
            top = min((l['y'] for l in ls if BL_TXT_RE.search(nospace(l['t']))), default=0)
            head_st = clean_status([q for a in an if a['rect'].y0 < top for q in a['lines'] if is_status_line(q)])
            if head_st:
                for r in recs.values():
                    if not r.get('status_raw'):
                        r['status_raw'] = head_st
                        r['pages'] = sorted(set(r['pages'] + [p]))      # หน้าที่มีคำสั่ง STATUS
            if kind == 'SHARED_CMS':
                # บล็อกคร่อมหน้า: สินค้าอยู่ท้ายหน้าก่อน แต่ D/O No. (เลข B/L) อยู่หัวหน้าถัดไป — ใส่หน้าก่อนด้วย
                for r in recs.values():
                    q = min(r['pages'])
                    ql = L[q - 1]
                    do_y = next((l['y'] for l in ql if r['bl'] in nospace(l['t'])), None)
                    if q > pages[0] and do_y is not None and not any('KGS' in l['t'] and l['y'] < do_y for l in ql):
                        r['pages'] = sorted(set(r['pages'] + [q - 1]))
            put(recs)
            continue
        if kind != 'TABLE':
            table_doc = None
        if kind == 'BLFORM':
            put(parse_blform(ls, an, p))
        elif kind == 'RIDER':
            bl, marks, typed = parse_rider(ls, p)
            if bl:
                ENT.setdefault(bl, blank(bl, p))
                r = ENT[bl]
                r['pages'] = sorted(set(r['pages'] + [p]))
                if marks:
                    r['rider_marks'] = marks
                if typed and not r['status_raw']:
                    r['status_raw'] = clean_status(typed)
        elif kind == 'MARKNO':
            put(parse_markno(ls, p))
        elif kind == 'TABLE':
            is_head = any('WE WOULD LIKE TO REQUEST' in l['t'].upper() or l['t'].upper() == 'AMENDMENT' for l in ls)
            if is_head or table_doc is None:
                table_doc = {'rows': {}, 'status': [], 'conts': [], 'st_pages': [], 'ct_pages': []}
                table_docs.append(table_doc)
            rows, st, cs = parse_table_page(ls, an, p)
            table_doc['rows'].update(rows)
            table_doc['status'] += st
            table_doc['conts'] += [c for c in cs if c not in table_doc['conts']]
            if st:
                table_doc['st_pages'].append(p)
            if cs:
                table_doc['ct_pages'].append(p)
        elif kind == 'LETTER':
            for bl, rec in parse_letter(ls, an, p).items():
                LETTERS[bl] = rec
        elif kind == 'SCAN':
            scans.append({'page': p, 'conts': conts_in(*[a['t'] for a in an])})

    for td in table_docs:
        for bl, rec in td['rows'].items():
            # STATUS/เลขตู้ที่ระบุครั้งเดียวทั้งฉบับ — ใส่หน้าที่มาด้วย เพื่อย้อนตรวจได้
            if not rec['status_raw'] and td['status']:
                rec['status_raw'] = ' '.join(dict.fromkeys(td['status']))
                rec['pages'] = sorted(set(rec['pages'] + td['st_pages']))
            if not rec['cont'] and td['conts']:
                rec['cont'] = ';'.join(td['conts'])
                rec['pages'] = sorted(set(rec['pages'] + td['ct_pages']))
            put({bl: finish(rec)})

    # ATTACHED RIDER: แทนคำ 'AS PER ATTACHED RIDER' ใน MARKS ด้วยเครื่องหมายจริงจาก RIDER
    for rec in ENT.values():
        rm = rec.pop('rider_marks', '')
        if rm:
            if re.search(r'AS\s*PER\s*ATTACHED\s*RIDER', rec['marks'], re.I):
                rec['marks'] = re.sub(r'AS\s*PER\s*ATTACHED\s*RIDER', rm, rec['marks'], flags=re.I).strip()
            elif not rec['marks']:
                rec['marks'] = rm

    # จดหมาย: B/L ที่มีฟอร์มโครงสร้างอยู่แล้ว ใช้จดหมายเติมเฉพาะช่องว่าง (STATUS ฯลฯ);
    # จดหมาย 'CONSIGNEE AS PER ATTACHED SHEET' ส่ง STATUS ให้ B/L ย่อยทุกตัวของ B/L นั้นด้วย
    for bl, lr in LETTERS.items():
        sibs = [b for b in ENT if b == bl or (lr.get('as_per_attached') and re.fullmatch(re.escape(bl) + r'[A-Z]?', b))]
        if sibs:
            for b in sibs:
                e = ENT[b]
                e['pages'] = sorted(set(e['pages'] + lr['pages']))
                if not e['status_raw'] and lr['status_raw']:
                    e['status_raw'] = lr['status_raw']
                if b == bl:
                    for k in ('cont', 'dg', 'reefer', 'transit_raw'):
                        if not e.get(k) and lr.get(k):
                            e[k] = lr[k]
        else:
            lr.pop('as_per_attached', None)
            ENT[bl] = lr
    for rec in ENT.values():
        rec.pop('as_per_attached', None)
    return ENT, info


def _fix_dsv(recs, doc):
    """DSV หน้ากว้าง (737pt) มีหลายตู้แบบ 'HLHU1234567 (40HC)' — layout G อ่านได้ตู้เดียว/ปนใน DESC"""
    for rec in recs.values():
        txt = ' '.join(doc[p - 1].get_text() for p in rec['pages'])
        cs = conts_in(txt)
        if cs:
            rec['cont'] = ';'.join(cs)
        rec['desc'] = ' '.join(w for w in re.split(r'(?<=\))\s+|\s{2,}', rec['desc'] or '')
                               if not re.fullmatch(r"[A-Z]{3}[UJZ]\d{7}\s*\(?[^)]*\)?", w.strip()))
        rec['desc'] = re.sub(r"\s*[A-Z]{3}[UJZ]\d{7}\s*\(\d{2}[A-Z]{0,2}\)", '', rec['desc']).strip()


# ═══════════════ กฎเฉพาะ HAHM ก่อนเทียบ ═══════════════
_ASPER_RE = re.compile(r'\*?\s*' + r'\s*'.join('ASPERATTACHED') + r'\s*(?:' + r'\s*'.join('SHEET') + '|'
                       + r'\s*'.join('RIDER') + r')?\s*\*?', re.I)   # 'AS PER ATTACHE D SHEET' (ช่องว่างกลางคำ)


def _subseq_extra(m_txt, e_txt):
    """MANIFEST เป็นส่วนหนึ่งของข้อความ B/L ใน ENTER (ตัวอักษร/ตัวเลขต่อกันตามลำดับ ไม่นับช่องว่าง)?
       คืนข้อความส่วนเกินของ ENTER หรือ None ถ้าไม่ใช่"""
    m_core = re.sub(r'[^A-Z0-9]', '', _ASPER_RE.sub(' ', str(m_txt or '').upper()))
    e_words = str(e_txt or '').split()
    if not m_core or not e_words:
        return None
    e_core = re.sub(r'[^A-Z0-9]', '', str(e_txt).upper())
    if m_core not in e_core:
        # ข้อความ MANIFEST อาจข้ามบรรทัดของ B/L (เช่นตัด 'HS CODE' ตรงกลาง) — ยอมให้เป็นลำดับย่อยทีละคำ
        it = iter(re.findall(r'[A-Z0-9]+', str(e_txt).upper()))
        if not all(any(w == x for x in it) for w in re.findall(r'[A-Z0-9]+', _ASPER_RE.sub(' ', str(m_txt).upper()))):
            return None
    extra = []
    rest = m_core
    for w in e_words:
        c = re.sub(r'[^A-Z0-9]', '', w.upper())
        if c and rest.startswith(c):
            rest = rest[len(c):]
        elif c:
            extra.append(w)
    return ' '.join(extra)


def apply_hahm_rules(MAN, ENT):
    for bl, m in MAN.items():
        # MANIFEST 'TEMP : - 18'C' (มีช่องว่างหลังเครื่องหมายลบ) — find_temp ของ skill อ่านไม่ออก
        if m.get('reefer') and '°C' not in m['reefer'] or (not m.get('reefer') and REEFER_RE.search(m.get('desc') or '')):
            m['reefer'] = find_reefer(m.get('reefer') or '', m.get('desc') or '') or m.get('reefer') or ''
    for bl, e in ENT.items():
        if re.search(r'\bREELS?\b', e.get('pkgtype') or '', re.I):
            e['pkgtype'] = 'REEL (ROLL)'
        m = MAN.get(bl)
        if not m or e.get('src') != 'BLFORM':
            continue
        # ARRIVAL NOTICE/สำเนา B/L ของ Heung-A เอง: MANIFEST ตัดข้อความ B/L บางส่วนออก (HS CODE, ที่อยู่ '*…',
        # TEL ฯลฯ) — ถ้า MANIFEST เป็นส่วนหนึ่งของ ENTER ตามลำดับ ถือว่าตรง แต่บอกส่วนเกินไว้ในหมายเหตุ
        for k in ('marks', 'desc'):
            if be.same_detail(e[k], m[k]) is False:
                extra = _subseq_extra(m[k], e[k])
                if extra is not None:
                    e[k + '_full'] = e[k]
                    e[k + '_extra'] = extra
                    e[k] = m[k]


def use_shed_by_status(ENT, shed_lcl, shed_other):
    """SHED ตาม STATUS: LCL / LCL/CFS → shed_lcl, อื่น ๆ (CY) → shed_other.
       STATUS ที่ใช้ = คำสั่งใน ENTER ถ้ามี (ลูกค้าสั่งจริง) ไม่มีจึงใช้ STATUS ของ MANIFEST"""
    orig = be.compute_extra_checks

    def wrapped(m, e):
        st = be.canon_status((e or {}).get('status_raw') or '') or be.canon_status(m.get('status') or '')
        be.EXPECTED_SHED = shed_lcl if st in ('LCL', 'LCL/CFS') else shed_other
        return orig(m, e)
    be.compute_extra_checks = wrapped


# ═══════════════ เติมเลขหน้า ENTER / หน้าสแกน ลงรายงาน ═══════════════
def annotate_report(outpath, MAN, ENT, info):
    wb = openpyxl.load_workbook(outpath)
    ws = wb.active
    head = [c.value for c in ws[4]]
    c_bl = head.index('B/L NO.') + 1
    c_note = len(head)
    c_full = {'marks': head.index('MARKS (ENTER)') + 1, 'desc': head.index('DESC. (ENTER)') + 1}
    c_sum = 1
    f_org = PatternFill('solid', fgColor='FCE4D6')
    scan_pages = [s['page'] for s in info['scans']]
    by_cont = {}
    for s in info['scans']:
        for c in s['conts']:
            for bl, m in MAN.items():
                if c in m['cont']:
                    by_cont.setdefault(bl, []).append(s['page'])
    for r in range(5, ws.max_row + 1):
        mbl = BL_TXT_RE.search(str(ws.cell(r, c_bl).value or ''))   # แถวเตือนขึ้นต้นด้วย '⚠ '
        if not mbl:
            continue
        bl = mbl.group(0)
        note = ws.cell(r, c_note)
        prefix = ''
        if bl in ENT and ENT[bl].get('pages'):
            pg = ENT[bl]['pages']
            prefix = f"[ENTER หน้า {', '.join(map(str, pg))}] "
            if ENT[bl].get('cons_note'):
                prefix += ENT[bl]['cons_note'] + ' ; '
            for k, lbl in (('marks', 'MARKS'), ('desc', 'DESC.')):
                if ENT[bl].get(k + '_full'):
                    ws.cell(r, c_full[k]).value = ENT[bl][k + '_full']
                    if ENT[bl].get(k + '_extra'):
                        prefix += f"{lbl} ตรง (MANIFEST เป็นส่วนหนึ่งของ B/L) — ENTER มีเพิ่ม: {ENT[bl][k + '_extra']} ; "
        elif bl in MAN and bl not in ENT:
            if bl in by_cont:
                pg = by_cont[bl]
                msg = (f"ENTER น่าจะเป็นหน้า {', '.join(map(str, pg))} (ภาพสแกน อ่านอัตโนมัติไม่ได้ — เลขตู้ที่พิมพ์ทับตรงกับ "
                       f"MANIFEST) — ตรวจด้วยตา")
            elif scan_pages:
                msg = (f"ไม่พบในข้อความ ENTER — มีหน้าภาพสแกนที่อ่านอัตโนมัติไม่ได้: หน้า {', '.join(map(str, scan_pages))} "
                       f"— ตรวจด้วยตา")
            else:
                msg = ''
            if msg:
                prefix = f'[{msg}] '
                for c in range(1, c_note + 1):
                    ws.cell(r, c).fill = f_org
                s = ws.cell(r, c_sum)
                s.value = f"⚠ ENTER เป็นภาพสแกน — ตรวจด้วยตา" if bl in by_cont else s.value
        if prefix:
            note.value = prefix + (note.value or '')
    wb.save(outpath)


# ═══════════════ main ═══════════════
def _find_one(patterns, folder):
    for pat in patterns:
        hits = glob.glob(os.path.join(folder, pat))
        if hits:
            return hits[0]
    return None


def main(argv=None):
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(encoding='utf-8', errors='replace')
        except Exception:
            pass
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--manifest', default=None)
    ap.add_argument('--enter', default=None)
    ap.add_argument('--indir', default='.', help='โฟลเดอร์งาน (มี input/) — ค่าเริ่มต้น: โฟลเดอร์ปัจจุบัน')
    ap.add_argument('--outdir', default=None)
    ap.add_argument('--shed', default='', help='SHED NO. เดียวทุก B/L เช่น 0141')
    ap.add_argument('--shed-lcl', default='', help='SHED ของ B/L ที่ STATUS เป็น CFS/LCL เช่น 0141')
    ap.add_argument('--shed-other', default='', help='SHED ของ B/L ที่เหลือ (CY) เช่น 0122')
    ap.add_argument('--name', default='EDI.xlsx', help='ชื่อไฟล์รายงาน (ค่าเริ่มต้น EDI.xlsx)')
    a = ap.parse_args(argv)
    _s = lambda v: (lambda m: m.group(0).zfill(4) if m else '')(re.search(r'\d{3,4}', v or ''))
    be.EXPECTED_SHED = _s(a.shed)
    if _s(a.shed_lcl) or _s(a.shed_other):
        use_shed_by_status(None, _s(a.shed_lcl), _s(a.shed_other))

    dirs = [os.path.join(a.indir, 'input'), a.indir]
    manifest = a.manifest or next((_find_one(['MANIFEST*.xls', 'MANIFEST*.xlsx', '*.xls', '*.xlsx'], d)
                                   for d in dirs if _find_one(['MANIFEST*.xls', 'MANIFEST*.xlsx', '*.xls', '*.xlsx'], d)), None)
    enter = a.enter or next((_find_one(['ENTER*.pdf', '*.pdf'], d) for d in dirs if _find_one(['ENTER*.pdf', '*.pdf'], d)), None)
    if not manifest or not enter:
        sys.exit('!! ไม่พบไฟล์ MANIFEST (.xls) หรือ ENTER (.pdf) — วางไว้ใน input/ หรือระบุ --manifest/--enter')
    outdir = a.outdir or os.path.join(a.indir, 'output')
    os.makedirs(outdir, exist_ok=True)
    outpath = os.path.join(outdir, a.name)
    if os.path.exists(outpath):
        sys.exit(f'!! มีไฟล์ {outpath} อยู่แล้ว — ไม่เขียนทับ ใช้ --name ชื่อใหม่')
    print(f'MANIFEST : {manifest}')
    print(f'ENTER    : {enter}')

    MAN, vessel, man_declared, bl_re = be.parse_manifest(manifest)
    ENT, info = parse_enter_hahm(enter, bl_re)

    # MANIFEST พิมพ์จำนวน/ชนิดหีบห่อซ้ำในวงเล็บต้น DESCRIPTION ('(5 WOODEN CASE) ...') — เหมือน build_edi.main()
    _n = lambda s: re.sub(r'[^A-Z0-9]', '', str(s).upper())
    for bl, e in ENT.items():
        mm = re.match(r'^(\([^()]*\))\s*', (MAN.get(bl) or {}).get('desc') or '')
        if mm and e.get('pkgs') and _n(mm.group(1)) == _n(f"{e['pkgs']}{e.get('pkgtype', '')}") \
                and not (e.get('desc') or '').startswith('('):
            e['desc'] = f"{mm.group(1)} {e.get('desc') or ''}".strip()

    apply_hahm_rules(MAN, ENT)
    stats = be.build_report(MAN, ENT, vessel, man_declared, {}, outpath)
    annotate_report(outpath, MAN, ENT, info)

    kinds = info['kinds']
    print('=' * 70)
    print('ชนิดหน้า ENTER:', ', '.join(f'{k}×{kinds.count(k)}' for k in dict.fromkeys(kinds)))
    if info['scans']:
        print('หน้าภาพสแกน (อ่านอัตโนมัติไม่ได้):', ', '.join(str(s['page']) for s in info['scans']))
    print(f'บันทึกแล้ว: {outpath}')
    print(f'B/L ทั้งหมด {stats["bl"]} | จุดต้องตรวจ {stats["issue_rows"]} แถว | STATUS ไม่ตรง {stats["st_bad"]} | '
          f'DG {stats["dg"]} | REEFER {stats["reefer"]} | TRANSIT {stats["transit"]}')
    print(f'SHED NO. ไม่ตรง {stats["shed_bad"]} | ประเทศปลายทาง ไม่ตรง {stats["dest_bad"]} | '
          f'CARGO MOVEMENT ไม่ตรง {stats["mv_bad"]} | TAX ID ไม่ตรง NOTIFY {stats["tax_bad"]}')
    missing = [b for b in MAN if b not in ENT]
    if missing:
        print(f'⚠️  B/L ที่ไม่มีข้อความ ENTER ({len(missing)}): {", ".join(missing)}')
    extra = [b for b in ENT if b not in MAN]
    if extra:
        print(f'⚠️  B/L ที่มีใน ENTER แต่ไม่มีใน MANIFEST ({len(extra)}): {", ".join(extra)}')
    if stats['crit']:
        print('\nจุดที่ต้องตรวจสอบ:')
        for x in stats['crit']:
            print('  -', x)
    print('=' * 70)


if __name__ == '__main__':
    main()
